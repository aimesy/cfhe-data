"""Read verified fields from the 2025 HCD Table B workbook layout.

visibility: public

This imports a submitted workbook, not the blank statewide template. It reads
saved values without executing macros or recalculating the source workbook.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import io
import math
import re
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .normalize import normalized_key


class TableBError(ValueError):
    """The source does not establish the requested Table B figures."""


def _jurisdiction_key(value: object) -> str:
    key = normalized_key(value)
    return re.sub(
        r"\s*(?:-\s*UNINCORPORATED(?: AREAS?)?|"
        r"\(\s*UNINCORPORATED(?: AREAS?)?\s*\))$",
        "",
        key,
    ).strip()


def _integer(value: object, cell: str, *, blank_zero: bool = False) -> int:
    if blank_zero and (value is None or str(value).strip() in {"", "-"}):
        return 0
    if isinstance(value, bool):
        raise TableBError(f"Table B!{cell}: boolean is not a count")
    if isinstance(value, str) and re.fullmatch(
        r"\d+|\d{1,3}(?:,\d{3})+", value.strip()
    ):
        value = int(value.strip().replace(",", ""))
    if not isinstance(value, (int, float)) or not math.isfinite(value):
        raise TableBError(f"Table B!{cell}: missing or invalid count {value!r}")
    if value < 0 or int(value) != value:
        raise TableBError(f"Table B!{cell}: count must be a nonnegative integer")
    return int(value)


def read_table_b_workbook(
    path: Path | str,
    expected_jurisdiction_key: str,
    expected_cycle: str,
    expected_reporting_year: int,
) -> dict[str, object]:
    """Validate identity, layout and arithmetic, then return four RHNA bands.

    Formula caches are required in counted cells. A genuinely empty detail
    cell or explicit dash is zero; a formula with no saved result is an error.
    Separate acutely/extremely low allocations require a different importer.
    """
    path = Path(path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise TableBError("Table B source must be an .xlsx or .xlsm workbook")
    if expected_reporting_year != 2025:
        raise TableBError("Only the verified 2025 Table B layout is supported")
    if expected_cycle not in {"6th", "7th"}:
        raise TableBError("Expected cycle must be 6th or 7th")
    source = path.read_bytes()
    grids = []
    for data_only in (False, True):
        try:
            workbook = load_workbook(
                io.BytesIO(source), read_only=True, data_only=data_only
            )
        except (BadZipFile, KeyError) as error:
            raise TableBError("Source is not a readable HCD workbook") from error
        try:
            if "Table B" not in workbook.sheetnames:
                raise TableBError("Workbook has no Table B sheet")
            grids.append(
                {
                    cell.coordinate: (cell.value, cell.data_type)
                    for row in workbook["Table B"].iter_rows(max_row=30, max_col=15)
                    for cell in row
                    if cell.value is not None
                }
            )
        finally:
            workbook.close()
    formulas, cached = grids

    def value(cell: str) -> object:
        result = cached.get(cell, (None, None))[0]
        if formulas.get(cell, (None, None))[1] == "f" and result is None:
            raise TableBError(f"Table B!{cell}: missing formula cache")
        return result

    def count(cell: str, *, blank_zero: bool = False) -> int:
        return _integer(value(cell), cell, blank_zero=blank_zero)

    labels = {
        "A1": "Jurisdiction",
        "A2": "Reporting Year",
        "A3": "Planning Period",
        "A7": "Table B",
        "A8": "Regional Housing Needs Allocation Progress",
        "A9": "Permitted Units Issued by Affordability",
        "D10": "Projection Period",
        "C11": "RHNA Allocation by Income Level",
        "N11": "Total Units to Date (all years)",
        "A13": "Acutely Low",
        "A15": "Extremely Low",
        "A17": "Very Low",
        "A19": "Low",
        "A21": "Moderate",
        "A23": "Above Moderate",
        "A24": "Total RHNA",
        "A25": "Total Units",
    }
    for row in (13, 15, 17, 19, 21):
        labels[f"B{row}"] = "Deed Restricted"
        labels[f"B{row + 1}"] = "Non-Deed Restricted"
    for cell, label in labels.items():
        if normalized_key(value(cell)) != normalized_key(label):
            raise TableBError(f"Table B!{cell}: unsupported layout, expected {label!r}")
    key = _jurisdiction_key(value("B1"))
    if not key or key != normalized_key(expected_jurisdiction_key):
        raise TableBError("Table B jurisdiction does not match expected canonical key")
    if count("B2") != expected_reporting_year:
        raise TableBError("Table B reporting year does not match expected year")
    if normalized_key(value("B3")) != normalized_key(f"{expected_cycle} Cycle"):
        raise TableBError("Table B cycle does not match expected cycle")
    period = re.fullmatch(
        r"\s*(\d{2}/\d{2}/\d{4})\s*[-–]\s*(\d{2}/\d{2}/\d{4})\s*",
        str(value("C3")),
    )
    if not period:
        raise TableBError("Table B!C3: invalid planning period")
    try:
        start, end = (
            dt.date(int(v[6:]), int(v[:2]), int(v[3:5])) for v in period.groups()
        )
    except ValueError as exc:
        raise TableBError("Table B!C3: invalid planning dates") from exc
    if not start < end or not start.year <= expected_reporting_year <= end.year:
        raise TableBError("Table B planning period does not contain the reporting year")
    if start.year > 2024 or count("C13") or count("C15"):
        raise TableBError("Separate acutely/extremely low allocations are unsupported")
    years = list(range(start.year, end.year + 1))
    if len(years) > 9:
        raise TableBError("Table B planning period exceeds the supported year columns")
    annual_columns = {}
    for index, col in enumerate(range(5, 14)):
        letter = get_column_letter(col)
        header = value(f"{letter}11")
        if index < len(years):
            if _integer(header, f"{letter}11") != years[index]:
                raise TableBError("Table B annual columns do not match planning period")
            annual_columns[letter] = years[index]
        elif header is not None and str(header).strip():
            raise TableBError("Table B has unexpected annual columns")
    detail = {
        (letter, row): count(f"{letter}{row}", blank_zero=True)
        for letter in (get_column_letter(col) for col in range(4, 14))
        for row in range(13, 24)
    }
    for letter in (get_column_letter(col) for col in range(4, 14)):
        column_sum = sum(detail[letter, row] for row in range(13, 24))
        if column_sum != count(f"{letter}25", blank_zero=True):
            raise TableBError(f"Table B!{letter}25: annual/projection total mismatch")
        if (
            letter != "D"
            and column_sum
            and annual_columns.get(letter, 9999) > expected_reporting_year
        ):
            raise TableBError("Table B includes units after the reporting cutoff")
    bands = {
        "vli": (17, range(13, 19)),
        "li": (19, range(19, 21)),
        "mi": (21, range(21, 23)),
        "ami": (23, range(23, 24)),
    }
    permits, allocations = {}, {}
    for band, (row, detail_rows) in bands.items():
        permits[band] = count(f"N{row}")
        allocations[band] = count(f"C{row}")
        summed = sum(detail[letter, r] for letter, r in detail if r in detail_rows)
        if permits[band] != summed:
            raise TableBError(
                f"Table B!N{row}: band total differs from annual/projection sum"
            )
    for row in (13, 15):
        subtotal = sum(v for (_, r), v in detail.items() if r in (row, row + 1))
        if count(f"N{row}") != subtotal:
            raise TableBError(
                f"Table B!N{row}: acutely/extremely low subtotal mismatch"
            )
    total, total_allocation = count("N25"), count("C24")
    if sum(permits.values()) != total or sum(allocations.values()) != total_allocation:
        raise TableBError("Table B four-band total or allocation total mismatch")
    return {
        "jurisdiction_key": key,
        "cycle": expected_cycle,
        "reporting_year": expected_reporting_year,
        "data_through": min(dt.date(expected_reporting_year, 12, 31), end).isoformat(),
        "period_start": start.isoformat(),
        "period_end": end.isoformat(),
        "permit_values": permits,
        "allocation_values": allocations,
        "total": total,
        "total_allocation": total_allocation,
        "projection_credit": count("D25", blank_zero=True),
        "projection_credit_by_income": {
            band: sum(detail["D", row] for row in detail_rows)
            for band, (_, detail_rows) in bands.items()
        },
        "source_sha256": hashlib.sha256(source).hexdigest(),
        "source_cells": {
            "sheet": "Table B",
            "jurisdiction": "B1",
            "reporting_year": "B2",
            "cycle": "B3",
            "planning_period": "C3",
            "total": "N25",
            "total_allocation": "C24",
            "projection_credit": "D25",
            "permit_values": {b: f"N{r}" for b, (r, _) in bands.items()},
            "allocation_values": {b: f"C{r}" for b, (r, _) in bands.items()},
            "detail": "D13:M23",
            "annual_totals": "D25:M25",
        },
    }
