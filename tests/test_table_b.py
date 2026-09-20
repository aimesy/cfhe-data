"""Table B imports reject ambiguous identities, missing caches and bad arithmetic.

visibility: public
"""

import hashlib
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from cfhe_data.table_b import TableBError, read_table_b_workbook


@pytest.fixture
def workbook_path(tmp_path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Table B"
    labels = {
        "A1": "Jurisdiction",
        "B1": "Example City",
        "A2": "Reporting Year",
        "B2": 2025,
        "A3": "Planning Period",
        "B3": "6th Cycle",
        "C3": "01/31/2023 - 01/31/2031",
        "A7": "Table B",
        "A8": "Regional Housing Needs Allocation Progress",
        "A9": "Permitted Units Issued by Affordability",
        "D10": "Projection Period",
        "C11": "RHNA Allocation by Income Level",
        "N11": "Total Units to Date (all years)",
        "A13": "Acutely Low",
        "A15": "Extremely Low",
        "A17": "Very Low",
        "A19": "Low",
        "A21": "Moderate",
        "A23": "Above Moderate",
        "A24": "Total RHNA",
        "A25": "Total Units",
        "C13": 0,
        "C15": 0,
        "C17": 100,
        "C19": 100,
        "C21": 100,
        "C23": 300,
        "C24": 600,
        "N13": 0,
        "N15": 2,
        "N17": 5,
        "N19": 4,
        "N21": 0,
        "N23": 10,
        "N25": 19,
        "D15": 2,
        "D17": 3,
        "D25": 5,
        "E19": 4,
        "E25": 4,
        "G23": 10,
        "G25": 10,
    }
    for row in (13, 15, 17, 19, 21):
        labels[f"B{row}"] = "Deed Restricted"
        labels[f"B{row + 1}"] = "Non-Deed Restricted"
    for cell, value in labels.items():
        sheet[cell] = value
    for col, year in enumerate(range(2023, 2032), 5):
        sheet.cell(11, col, year)
    path = tmp_path / "submitted.xlsx"
    workbook.save(path)
    return path


def read(path: Path, **kwargs: object) -> dict[str, object]:
    return read_table_b_workbook(
        path,
        **{
            "expected_jurisdiction_key": "EXAMPLE CITY",
            "expected_cycle": "6th",
            "expected_reporting_year": 2025,
            **kwargs,
        },
    )


def change(path: Path, cell: str, value: object) -> None:
    workbook = load_workbook(path)
    workbook["Table B"][cell] = value
    workbook.save(path)


def test_includes_projection_and_folds_eli_once(workbook_path: Path) -> None:
    result = read(workbook_path)
    assert result["permit_values"] == {"vli": 5, "li": 4, "mi": 0, "ami": 10}
    assert result["total"] == 19
    assert result["projection_credit"] == 5
    assert result["projection_credit_by_income"] == {
        "vli": 5,
        "li": 0,
        "mi": 0,
        "ami": 0,
    }
    assert result["total_allocation"] == 600
    assert result["data_through"] == "2025-12-31"
    assert (
        result["source_sha256"]
        == hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    )


def test_projection_credit_retains_each_income_band(workbook_path: Path) -> None:
    workbook = load_workbook(workbook_path)
    for cell, value in {
        "D19": 7,
        "D21": 3,
        "D23": 11,
        "D25": 26,
        "N19": 11,
        "N21": 3,
        "N23": 21,
        "N25": 40,
    }.items():
        workbook["Table B"][cell] = value
    workbook.save(workbook_path)
    result = read(workbook_path)
    assert result["projection_credit_by_income"] == {
        "vli": 5,
        "li": 7,
        "mi": 3,
        "ami": 11,
    }
    assert result["projection_credit"] == 26


@pytest.mark.parametrize(
    "cell,value,message",
    [
        ("B1", "Other City", "jurisdiction"),
        ("B2", 2024, "reporting year"),
        ("B3", "7th Cycle", "cycle"),
        ("C3", "01/31/2031 - 01/31/2023", "planning period"),
        ("E11", 2022, "annual columns"),
        ("N11", "Current year", "unsupported layout"),
        ("N17", 7, "band total"),
        ("N25", 21, "four-band total"),
        ("C24", 601, "allocation total"),
        ("C15", 10, "Separate"),
        ("N17", None, "missing or invalid"),
        ("N17", 1.5, "integer"),
        ("N17", True, "boolean"),
        ("N17", -1, "nonnegative"),
        ("N17", "#VALUE!", "invalid"),
        ("D25", 7, "projection total"),
        ("D17", "=SUM(1,2)", "missing formula cache"),
        ("N17", "=SUM(D13:M18)", "missing formula cache"),
    ],
)
def test_rejects_unverified_source(
    workbook_path: Path, cell: str, value: object, message: str
) -> None:
    change(workbook_path, cell, value)
    with pytest.raises(TableBError, match=message):
        read(workbook_path)


def test_rejects_future_values(workbook_path: Path) -> None:
    change(workbook_path, "H23", 2)
    change(workbook_path, "H25", 2)
    with pytest.raises(TableBError, match="after the reporting cutoff"):
        read(workbook_path)


def test_narrow_county_alias(workbook_path: Path) -> None:
    change(workbook_path, "B1", "Example County - Unincorporated")
    assert (
        read(workbook_path, expected_jurisdiction_key="EXAMPLE COUNTY")[
            "jurisdiction_key"
        ]
        == "EXAMPLE COUNTY"
    )


def test_malformed_workbook_has_a_table_b_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not an Excel archive")
    with pytest.raises(TableBError, match="not a readable HCD workbook"):
        read(path)
